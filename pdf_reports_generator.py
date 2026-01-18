# --- FILE: pdf_reports_generator.py ---

import os
import pandas as pd
from fpdf import FPDF
from utils import resource_path # Or wherever you put the function
from docx.enum.text import WD_ALIGN_PARAGRAPH # Needed if manipulating docx tables
from docx.enum.table import WD_ALIGN_VERTICAL # Needed if manipulating docx tables
import logging
import time
# --- Constants ---
LOGO_PATH = resource_path('image_5eafd9.png') # Assumes the logo is in the same directory

# ... (função sanitize_text permanece igual) ...
def sanitize_text(text):
    """
    Encodes text to latin-1, replacing any unsupported characters to prevent crashes.
    """
    return str(text).encode('latin-1', 'replace').decode('latin-1')

# ... (função _safe_pdf_output permanece igual) ...
def _safe_pdf_output(pdf, path, retries=5, delay=0.5):
    for i in range(retries):
        try:
            pdf.output(path)
            logging.info(f"PDF Save: File saved successfully on attempt {i+1}: {path}")
            return
        except PermissionError as e:
            logging.warning(f"PDF Save: Attempt {i+1} failed. File is locked: {path}. Retrying in {delay}s...")
            time.sleep(delay)
        except Exception as e:
            if "cannot find" in str(e) or "font" in str(e):
                 logging.error(f"PDF Save: Critical FPDF font error: {e}")
            raise e
    logging.error(f"PDF Save: Failed to save after {retries} attempts: {path}. File remains locked.")
    raise PermissionError(f"Could not save PDF, file remains locked: {path}")

# ... (função safe_strftime permanece igual) ...
def safe_strftime(dt, fmt="%d/%m/%Y"):
    try:
        if pd.notna(dt):
            return dt.strftime(fmt)
    except Exception:
        pass
    return "N/A"

# ... (função _format_brl permanece igual) ...
def _format_brl(value):
    """Formats a number into BRL currency string for PDF."""
    try:
        return f"{value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    except (ValueError, TypeError):
        return "0,00" # Or handle as appropriate

# ... (função _prepare_general_analysis_data permanece igual) ...
def _prepare_general_analysis_data(df_all_invoices):
    logging.info("PDF Gen: Iniciando _prepare_general_analysis_data...")
    invoices_list = []
    
    if df_all_invoices is None or df_all_invoices.empty:
        logging.warning("PDF Gen: df_all_invoices está vazio. Retornando lista vazia.")
        return []

    df_sorted = df_all_invoices.copy()

    logging.info("PDF Gen: Verificando colunas essenciais para o PDF de Análise Geral...")
    if 'DATA EMISSÃO' not in df_sorted.columns:
        logging.warning("PDF Gen: Coluna 'DATA EMISSÃO' não encontrada. A adicionar coluna NaT.")
        df_sorted['DATA EMISSÃO'] = pd.NaT
    elif not pd.api.types.is_datetime64_any_dtype(df_sorted['DATA EMISSÃO']):
        df_sorted['DATA EMISSÃO'] = pd.to_datetime(df_sorted['DATA EMISSÃO'], errors='coerce')

    if 'VALOR' not in df_sorted.columns:
        logging.warning("PDF Gen: Coluna 'VALOR' não encontrada. A adicionar coluna com 0.0.")
        df_sorted['VALOR'] = 0.0
    if 'VALOR_ORIGINAL' not in df_sorted.columns:
        df_sorted['VALOR_ORIGINAL'] = df_sorted['VALOR'] # Fallback
    
    logging.info("PDF Gen: Colunas verificadas. A ordenar por 'DATA EMISSÃO'.")

    df_sorted = df_sorted.sort_values(by="DATA EMISSÃO")

    df_sorted['VALOR CONSIDERADO'] = df_sorted['VALOR'] - df_sorted.get('VALOR DEDUÇÃO', 0)
    logging.info("PDF Gen: Ordenação e cálculo de 'VALOR CONSIDERADO' concluídos.")

    for _, row in df_sorted.iterrows():
        invoices_list.append({
            'numero': sanitize_text(row.get('NÚMERO', '')),
            'data_emissao': safe_strftime(row.get('DATA EMISSÃO', pd.NaT)),
            'valor': _format_brl(row.get('VALOR_ORIGINAL', 0.0)),
            'desconto_incondicional': _format_brl(row.get('DESCONTO INCONDICIONAL', 0.0)),
            'valor_deducao': _format_brl(row.get('VALOR DEDUÇÃO', 0.0)),
            'valor_considerado': _format_brl(row.get('VALOR CONSIDERADO', 0.0)),
            'cnpj_tomador': sanitize_text(row.get('CNPJ/CPF TOMADOR', '')),
            'tomador': sanitize_text(row.get('TOMADOR', ''))
        })
    
    logging.info(f"PDF Gen: _prepare_general_analysis_data concluído. {len(invoices_list)} notas processadas.")
    return invoices_list

# ... (função _prepare_infraction_auto_data permanece igual) ...
def _prepare_infraction_auto_data(df_infractions_filtered):
    if df_infractions_filtered.empty:
        return [], {}

    df_copy = df_infractions_filtered.copy()
    if not pd.api.types.is_datetime64_any_dtype(df_copy['DATA EMISSÃO']):
        df_copy['DATA EMISSÃO'] = pd.to_datetime(df_copy['DATA EMISSÃO'], errors='coerce')

    if df_copy.empty or df_copy['DATA EMISSÃO'].isnull().all():
        return [], {}

    # ✅ CHANGE 1: Create Year and Month columns for grouping
    df_copy['ANO'] = df_copy['DATA EMISSÃO'].dt.year
    df_copy['MÊS'] = df_copy['DATA EMISSÃO'].dt.month
    
    # Sort chronologically (Year first, then Month)
    df_sorted = df_copy.sort_values(by=['ANO', 'MÊS', 'DATA EMISSÃO'])

    processed_rows = []
    grand_total_valor_original = 0
    grand_total_desconto = 0
    grand_total_base_calculo = 0

    # ✅ CHANGE 2: Group by BOTH Year and Month
    for (year, month), group in df_sorted.groupby(['ANO', 'MÊS']):
        month_total_valor_original = group.get('VALOR_ORIGINAL', 0.0).sum()
        month_total_desconto = group.get('DESCONTO INCONDICIONAL', 0.0).sum()
        month_total_base_calculo = group.get('VALOR', 0.0).sum()

        for _, row in group.iterrows():
            base_calculo = row.get('VALOR', 0.0)
            processed_rows.append({
                'is_subtotal': False,
                'numero': sanitize_text(row.get('NÚMERO', '')),
                'data_emissao': safe_strftime(row.get('DATA EMISSÃO', pd.NaT)),
                'valor': _format_brl(row.get('VALOR_ORIGINAL', 0.0)),
                'desconto_incondicional': _format_brl(row.get('DESCONTO INCONDICIONAL', 0.0)),
                'aliquota': f"{row.get('ALÍQUOTA', 0.0):.2f}%",
                'regime': sanitize_text(row.get('REGIME DE TRIBUTAÇÃO', '')),
                'natureza': sanitize_text(row.get('NATUREZA DA OPERAÇÃO', '')),
                'iss_retido': sanitize_text(row.get('ISS RETIDO', 'Não')),
                'discriminacao': sanitize_text(row.get('DISCRIMINAÇÃO DOS SERVIÇOS', '')),
                'codigo_atividade': sanitize_text(row.get('CÓDIGO DA ATIVIDADE', '')),
                'pagamento': sanitize_text(row.get('PAGAMENTO', '')),
                'mes': month, 
                'ano': year, # Pass year to row data
                'base_calculo': _format_brl(base_calculo),
                'aliquota_correta': f"{row.get('correct_rate', 5.0):.2f}%"
            })

        # ✅ CHANGE 3: Subtotal Label includes MM/YYYY
        processed_rows.append({
            'is_subtotal': True, 
            'mes': f"{month:02d}/{year}", 
            'total_valor': _format_brl(month_total_valor_original),
            'total_desconto': _format_brl(month_total_desconto),
            'total_base_calculo': _format_brl(month_total_base_calculo),
        })
        grand_total_valor_original += month_total_valor_original
        grand_total_desconto += month_total_desconto
        grand_total_base_calculo += month_total_base_calculo

    grand_total = {
        'valor': _format_brl(grand_total_valor_original),
        'desconto': _format_brl(grand_total_desconto),
        'base_calculo': _format_brl(grand_total_base_calculo),
    }
    return processed_rows, grand_total

# ... (classe PDFReport permanece igual) ...
class PDFReport(FPDF):
    def __init__(self, company_context, **kwargs):
        super().__init__(**kwargs)
        self.company_context = company_context
        self.set_auto_page_break(auto=True, margin=15)

    def header(self):
        logo_drawn_width = 0
        logo_padding = 0
        if os.path.exists(LOGO_PATH):
            self.image(LOGO_PATH, x=self.l_margin, y=8, w=25)
            logo_drawn_width = 25
            logo_padding = 5
        
        self.set_y(10)
        self.set_font('Helvetica', 'B', 10)
        self.cell(0, 5, 'PREFEITURA MUNICIPAL DE CURITIBA', 0, 1, 'C')
        self.cell(0, 5, 'SECRETARIA MUNICIPAL DE PLANEJAMENTO, FINANÇAS E ORÇAMENTO', 0, 1, 'C')
        self.cell(0, 5, 'DEPARTAMENTO DE RENDAS MOBILIÁRIAS', 0, 1, 'C')
        self.ln(5)
        
        y_pos_start = self.get_y()

        company_x_pos = self.l_margin 
        self.set_x(company_x_pos) 
        self.set_font('Helvetica', '', 9)
        company_info = (
            f"NOME/RAZÃO SOCIAL: {sanitize_text(self.company_context.get('razao_social', ''))}\n"
            f"IMU: {sanitize_text(self.company_context.get('imu', ''))}\n"
            f"CNPJ: {sanitize_text(self.company_context.get('cnpj', ''))}"
        )
        self.multi_cell(120, 5, company_info, 0, 'L')
        y_after_company = self.get_y() 

        self.set_y(y_pos_start) 
        address_width = 50
        address_x_pos = self.w - self.r_margin - address_width
        self.set_x(address_x_pos)
        self.set_font('Helvetica', '', 9)
        self.multi_cell(address_width, 5,
            "Av. Cândido de Abreu, nº 817 - Térreo\n"
            "Centro Cívico\n"
            "80530-908 Curitiba-Paraná", 0, 'R')
        y_after_address = self.get_y()

        final_y = max(y_after_company, y_after_address)
        self.set_y(final_y + 5)

# ... (função _draw_table_row permanece igual) ...
def _draw_table_row(pdf, data, col_widths, border=1, fill=False, align='C', font_style='', font_size=5):
    pdf.set_font('Helvetica', font_style, font_size)
    line_height = pdf.font_size * 1.5
    num_lines = []
    
    for i, datum in enumerate(data):
        lines = pdf.multi_cell(w=col_widths[i] - 2, h=line_height, txt=str(datum), border=0, split_only=True)
        num_lines.append(len(lines))
    
    max_lines = max(num_lines) if num_lines else 1
    total_row_height = (max_lines * line_height) + 2
    printable_page_height = pdf.h - pdf.b_margin
    if pdf.get_y() + total_row_height > printable_page_height:
        pdf.add_page(orientation=pdf.cur_orientation)
    
    start_y = pdf.get_y()
    start_x = pdf.get_x()
    current_x = start_x
    
    for i, datum in enumerate(data):
        align_char = align[i] if isinstance(align, list) else align
        if fill:
            pdf.set_fill_color(230, 230, 230)
            pdf.rect(current_x, start_y, col_widths[i], total_row_height, 'F')
        if border:
            pdf.rect(current_x, start_y, col_widths[i], total_row_height, 'D')
        
        text_padding_x = 1
        text_padding_y = 1
        pdf.set_xy(current_x + text_padding_x, start_y + text_padding_y)
        
        pdf.multi_cell(
            w=col_widths[i] - (text_padding_x * 2), h=line_height,
            txt=str(datum), # Usa str(datum) diretamente
            border=0, align=align_char, fill=False
        )
        current_x += col_widths[i]
    
    pdf.set_y(start_y + total_row_height)

# ... (função _create_general_analysis_pdf permanece igual) ...
def _create_general_analysis_pdf(company_context, all_invoices, output_path):
    pdf = PDFReport(company_context, orientation='L', unit='mm', format='A4')
    pdf.add_page()

    pdf.set_font('Helvetica', 'B', 12)
    pdf.cell(0, 10, 'RELATORIO DE NOTAS FISCAIS ANALISADAS NA OPERACAO', 0, 1, 'C') # ✅ Texto sanitizado
    pdf.ln(5)

    col_widths = [15, 25, 25, 25, 25, 25, 45, 85]
    headers = ['NÚMERO', 'DATA EMISSÃO', 'VALOR\n(Original)', 'DESCONTO\nINCONDIC.', 'VALOR\nDEDUÇÃO', 'VALOR\nCONSIDERADO', 'CNPJ/CPF\nTOMADOR', 'TOMADOR']

    _draw_table_row(pdf, headers, col_widths, fill=True, font_style='B', font_size=5)

    for invoice in all_invoices:
        row_data = [
            invoice['numero'], invoice['data_emissao'], invoice['valor'],
            invoice['desconto_incondicional'], invoice['valor_deducao'],
            invoice['valor_considerado'], invoice['cnpj_tomador'], invoice['tomador']
        ]
        alignments = ['C', 'C', 'R', 'R', 'R', 'R', 'L', 'L']
        _draw_table_row(pdf, row_data, col_widths, align=alignments, font_size=5)
    
    logging.info("PDF Gen: Todas as linhas do PDF de Análise Geral foram desenhadas.")
    logging.info(f"PDF Gen: A tentar salvar o PDF em: {output_path}")
    try:
        _safe_pdf_output(pdf, output_path)
        logging.info("PDF Gen: PDF de Análise Geral salvo com sucesso.")
    except Exception as e:
        logging.error(f"PDF Gen: CRASH ao salvar o PDF: {output_path}")
        logging.exception(e)
        raise e

# ... (função _create_infraction_auto_pdf permanece igual) ...
def _create_infraction_auto_pdf(company_context, auto_id, invoice_rows, grand_total, is_compensated_auto, output_path):
    pdf = PDFReport(company_context, orientation='L', unit='mm', format='A4')
    pdf.add_page()

    pdf.set_font('Helvetica', 'B', 10)
    pdf.cell(0, 5, sanitize_text(f"AUTO DE INFRAÇÃO Nº: {auto_id}"), 0, 1, 'L')
    pdf.ln(5)

    # 14 Columns defined here
    col_widths = [10, 16, 12, 12, 10, 25, 20, 12, 45, 18, 14, 8, 18, 18]
    headers = ['NÚMERO', 'DATA\nEMISSÃO', 'VALOR\n(Original)', 'DESC.\nINCOND.', 'ALÍQUO\nTA',
               'REGIME DE\nTRIBUTAÇÃO', 'NATUREZA DA\nOPERAÇÃO', 'ISS\nRETIDO',
               'DISCRIMINAÇÃO DOS SERVIÇOS', 'CÓDIGO DA\nATIVIDADE', 'PAGAMEN\nTO', 'MÊS',
               'BASE DE\nCÁLCULO', 'ALÍQUOTA\nCORRETA']

    _draw_table_row(pdf, headers, col_widths, fill=True, font_style='B', font_size=5)

    for row in invoice_rows:
        if row['is_subtotal']:
            # ... (Subtotal logic remains valid) ...
            row_height = 8
            pdf.set_font('Helvetica', 'B', 6.5)
            pdf.set_fill_color(230, 230, 230)
            pdf.cell(sum(col_widths[:2]), row_height, f"Total {row['mes']}", 1, 0, 'C', 1)
            pdf.cell(col_widths[2], row_height, row['total_valor'], 1, 0, 'R', 1)
            pdf.cell(col_widths[3], row_height, row['total_desconto'], 1, 0, 'R', 1)
            pdf.cell(sum(col_widths[4:12]), row_height, '', 1, 0, 'C', 1)
            pdf.cell(col_widths[12], row_height, row['total_base_calculo'], 1, 0, 'R', 1)
            pdf.cell(col_widths[13], row_height, '', 1, 1, 'C', 1)
        else:
            # ✅ FIX: Removed the extra "str(row['mes'])" that caused the index error.
            # Now 'data' has 14 items, matching 'col_widths' (14 items).
            data = [
                row['numero'], 
                row['data_emissao'], 
                row['valor'], 
                row['desconto_incondicional'], 
                row['aliquota'],
                row['regime'], 
                row['natureza'], 
                row['iss_retido'],
                row['discriminacao'], 
                row['codigo_atividade'], 
                row['pagamento'],
                f"{row['mes']:02d}/{str(row['ano'])[-2:]}", # This is the MÊS column
                # ❌ REMOVED: str(row['mes']), <--- This was the 15th item causing the crash
                row['base_calculo'], 
                row['aliquota_correta']
            ]
            
            aligns = ['C', 'C', 'R', 'R', 'C', 'L', 'L', 'C', 'L', 'C', 'C', 'C', 'R', 'C']
            _draw_table_row(pdf, data, col_widths, align=aligns, font_size=5)

    # ... (Rest of function remains the same) ...
    # Grand Total Row
    pdf.set_font('Helvetica', 'B', 7)
    total_row_height = 7
    pdf.set_fill_color(230, 230, 230)
    pdf.cell(sum(col_widths[:2]), total_row_height, "Total Geral", 1, 0, 'C', 1)
    pdf.cell(col_widths[2], total_row_height, grand_total.get('valor', '0,00'), 1, 0, 'R', 1)
    pdf.cell(col_widths[3], total_row_height, grand_total.get('desconto', '0,00'), 1, 0, 'R', 1)
    pdf.cell(sum(col_widths[4:12]), total_row_height, '', 1, 0, 'C', 1)
    pdf.cell(col_widths[12], total_row_height, grand_total.get('base_calculo', '0,00'), 1, 0, 'R', 1)
    pdf.cell(col_widths[13], total_row_height, '', 1, 1, 'C', 1)

    if is_compensated_auto:
        pdf.ln(5)
        pdf.set_font('Helvetica', 'I', 8)
        pdf.cell(0, 5, sanitize_text("Nota: O valor total deste auto de infração foi totalmente compensado por pagamentos (DAM/DAS) efetuados."), 0, 1, 'L')

    _safe_pdf_output(pdf, output_path)

# --- Main Generation Function ---
def generate_detailed_pdfs(company_context, all_invoices_df, final_data, preview_context, output_dir, status_callback=None):
    
    # ✅ --- INÍCIO DA CORREÇÃO ---
    # 'status_callback' agora é a função 'emit' do 'generation_task'
    # Não é um objeto Qt, por isso chamamo-lo diretamente.
    if status_callback:
        emit = status_callback
    else:
        emit = print
    # ✅ --- FIM DA CORREÇÃO ---
    
    logging.info("PDF Gen: Iniciando generate_detailed_pdfs...")

    file_prefix = company_context.get('cnpj', 'unknown').replace('/', '').replace('.', '').replace('-', '')

    try:
        emit("📄 Gerando PDF de Análise Geral de Notas...")
        logging.info("PDF Gen: A chamar _prepare_general_analysis_data...")
        all_invoices_data = _prepare_general_analysis_data(all_invoices_df)
        
        logging.info("PDF Gen: _prepare_general_analysis_data concluído. A definir caminho de saída.")
        output_path = os.path.join(output_dir, "notas_analise_geral.pdf")
        
        logging.info(f"PDF Gen: A chamar _create_general_analysis_pdf para: {output_path}")
        _create_general_analysis_pdf(company_context, all_invoices_data, output_path)
        
        logging.info("PDF Gen: PDF de Análise Geral gerado.")
        emit("✅ PDF de Análise Geral gerado com sucesso.")
    except Exception as e:
        logging.error(f"PDF Gen: CRASH ao gerar PDF de Análise Geral: {e}")
        logging.exception(e) # Loga o traceback completo
        emit(f"❌ ERRO ao gerar PDF de Análise Geral: {e}")

    try:
        emit("⚖️ Gerando PDFs detalhados por Auto de Infração...")
        logging.info("PDF Gen: A gerar PDFs por Auto de Infração...")

        if not preview_context or 'autos' not in preview_context:
            logging.warning("PDF Gen: 'preview_context' está em falta ou vazio.")
            emit("   - AVISO: 'preview_context' está em falta ou vazio. Não é possível determinar os totais dos autos.")
            return

        all_autos_preview_map = {
            auto.get('numero'): auto 
            for auto in preview_context.get('autos', []) 
            if auto.get('numero') is not None
        }

        for auto_id, auto_info_final_data in final_data.items():
            invoice_indices = auto_info_final_data.get('invoices', [])
            if not invoice_indices:
                emit(f"   - AVISO: Pulando PDF para o Auto '{auto_id}' por não ter faturas associadas (final_data).")
                continue

            valid_indices = [idx for idx in invoice_indices if idx in all_invoices_df.index]
            if not valid_indices:
                emit(f"   - AVISO: Pulando PDF para o Auto '{auto_id}' - índices de fatura inválidos ou não encontrados.")
                continue

            df_auto_invoices_original = all_invoices_df.loc[valid_indices].copy()
            auto_data_preview = all_autos_preview_map.get(auto_id)
            is_compensated_auto = False
            
            df_invoices_para_pdf = df_auto_invoices_original.copy()

            if auto_data_preview:
                if auto_data_preview.get('totais', {}).get('iss_apurado_op', 0.0) <= 0.01:
                    is_compensated_auto = True
                else:
                    dados_anuais = auto_data_preview.get('dados_anuais', [])
                    meses_compensados_set = set()
                    
                    for mes_data in dados_anuais:
                        iss_val = mes_data.get('iss_apurado_op', 0.0)
                        mes_ano_val = mes_data.get('mes_ano')
                        
                        if iss_val <= 0.01 and mes_ano_val:
                            meses_compensados_set.add(mes_ano_val)

                    if meses_compensados_set and 'DATA EMISSÃO' in df_invoices_para_pdf.columns:
                        try:
                            df_invoices_para_pdf['mes_ano_str'] = pd.to_datetime(df_invoices_para_pdf['DATA EMISSÃO'], errors='coerce').dt.strftime('%m/%Y')
                            df_invoices_para_pdf = df_invoices_para_pdf[~df_invoices_para_pdf['mes_ano_str'].isin(meses_compensados_set)]
                        except Exception as e:
                            emit(f"   - ERRO ao filtrar meses compensados para Auto '{auto_id}': {e}")
                    elif meses_compensados_set:
                        emit(f"   - AVISO: Não foi possível filtrar meses compensados para Auto '{auto_id}' - coluna 'DATA EMISSÃO' não encontrada.")
            else:
                emit(f"   - AVISO: Dados de preview não encontrados para Auto '{auto_id}'. Não é possível determinar compensação.")

            processed_rows, grand_total = _prepare_infraction_auto_data(df_invoices_para_pdf)

            auto_name_sanitized = auto_id.replace(" ", "_").replace(":", "").replace("%", "").replace(",", "").replace("/", "")
            output_path = os.path.join(output_dir, f"notas_auto_{auto_name_sanitized}.pdf")

            _create_infraction_auto_pdf(company_context, auto_id, processed_rows, grand_total, is_compensated_auto, output_path)

            emit(f"   - PDF para o Auto '{auto_id}' gerado.")
        emit("✅ PDFs por Auto de Infração gerados com sucesso.")
    except Exception as e:
        emit(f"❌ ERRO ao gerar PDFs por Auto de Infração: {e}")